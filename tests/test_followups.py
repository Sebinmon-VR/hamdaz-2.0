"""When a task is asked about, and what the two emails say.

No database, no SharePoint, no mail: the rule is a function of one task, a
clock and the settings, and the emails are functions of one row.
"""

from __future__ import annotations

from datetime import UTC, datetime, timedelta

from app.followups import mailer
from app.followups.service import decide, due_of
from app.models.followup import FollowupStatus, TaskFollowup
from app.models.user import User
from app.proposals.sharepoint import ProposalTask

DUE = datetime(2026, 9, 28, 10, 0, tzinfo=UTC)
WATCH = datetime(2026, 9, 28, 9, 0, tzinfo=UTC)


def stored(meant: datetime) -> str:
    """How SharePoint stores a time typed in the UAE: the site is on US Pacific
    time, so the typed clock is taken as Pacific and kept in UTC."""
    from zoneinfo import ZoneInfo

    wall = meant.astimezone(ZoneInfo("Asia/Dubai")).replace(tzinfo=None)
    return wall.replace(tzinfo=ZoneInfo("America/Los_Angeles")).astimezone(UTC).isoformat().replace("+00:00", "Z")


def task(**overrides) -> ProposalTask:
    fields = dict(
        id="901", title="test — overdue follow-up", status="In Progress", priority=None,
        assigned_to_lookup_id="12", assigned_to_name="Sebin", start_date=None,
        due_date=None, bid_closing_date=stored(DUE),
        end_user="ADNOC", submission_status=None, current_type=None, order_status=None,
        negotiation=None, quote_no=None, remarks=None, working_notes=None,
        created_at=None, modified_at=None,
    )
    fields.update(overrides)
    return ProposalTask(**fields)


def at(minutes_after_due: int) -> datetime:
    return DUE + timedelta(minutes=minutes_after_due)


def test_asked_once_the_grace_after_the_due_time_has_passed() -> None:
    assert not decide(task(), now=at(19), grace_minutes=20, watch_from=WATCH).ask
    ruled = decide(task(), now=at(20), grace_minutes=20, watch_from=WATCH)
    assert ruled.ask and ruled.due_at == DUE


def test_a_submitted_bid_is_never_asked_about() -> None:
    ruled = decide(task(submission_status="Submitted"), now=at(60), grace_minutes=20, watch_from=WATCH)
    assert not ruled.ask and ruled.why == "finished"


def test_completed_but_not_submitted_is_still_asked_about() -> None:
    """The case from the list: Status says Completed, Submission Status says
    Not Submitted. The submission decides."""
    both = task(status="Completed", submission_status="Not Submitted")
    assert decide(both, now=at(60), grace_minutes=20, watch_from=WATCH).ask
    assert decide(task(status="Completed"), now=at(60), grace_minutes=20, watch_from=WATCH).ask


def test_marked_not_submitted_is_asked_at_once_not_at_the_deadline() -> None:
    marked = task(status="Completed", submission_status="Not Submitted")
    ruled = decide(marked, now=at(-300), grace_minutes=20, watch_from=WATCH)
    assert ruled.ask and ruled.why == "marked not submitted"
    # A blank submission status is still work in hand: it waits for the due time.
    assert not decide(task(), now=at(-300), grace_minutes=20, watch_from=WATCH).ask


def test_a_task_with_no_status_is_asked_about_because_that_is_the_point() -> None:
    assert decide(task(status=None), now=at(60), grace_minutes=20, watch_from=WATCH).ask


def test_nothing_due_before_the_watch_began_is_asked_about() -> None:
    old = task(bid_closing_date=stored(datetime(2025, 11, 27, tzinfo=UTC)))
    ruled = decide(old, now=at(60), grace_minutes=20, watch_from=WATCH)
    assert not ruled.ask and ruled.why == "due before the watch began"
    assert not decide(task(), now=at(60), grace_minutes=20, watch_from=None).ask


def test_the_title_filter_keeps_a_trial_to_the_test_tasks() -> None:
    real = task(title="RFQ 6000129780 Customization")
    assert not decide(real, now=at(60), grace_minutes=20, watch_from=WATCH, title_contains="test").ask
    assert decide(task(), now=at(60), grace_minutes=20, watch_from=WATCH, title_contains="TEST").ask


def test_the_bid_closing_time_is_read_as_it_was_typed() -> None:
    """The case from the list: SharePoint shows BCD UAE Time 8:29 AM and holds
    15:29 UTC, because the site is on Pacific time. The bid closes at 8:29 AM
    in the UAE — 04:29 UTC — not at 19:29 as a plain UTC reading had it."""
    rameesa = task(bid_closing_date="2026-09-28T15:29:00Z")
    assert due_of(rameesa) == datetime(2026, 9, 28, 4, 29, tzinfo=UTC)
    # And the two that were missing from today: 1:00 PM and 1:30 PM UAE.
    assert due_of(task(bid_closing_date="2026-09-28T20:00:00Z")) == datetime(2026, 9, 28, 9, 0, tzinfo=UTC)
    assert due_of(task(bid_closing_date="2026-09-28T20:30:00Z")) == datetime(2026, 9, 28, 9, 30, tzinfo=UTC)


def test_the_bid_closing_time_leads_and_a_due_date_means_the_end_of_that_day() -> None:
    both = task(due_date="2026-09-27T07:00:00Z", bid_closing_date="2026-09-28T13:51:46Z")
    assert due_of(both) == datetime(2026, 9, 28, 2, 51, 46, tzinfo=UTC)
    # A date-only Due Date is midnight on the site's clock (07:00Z in Pacific
    # summer time); it is due at 23:59:59 UAE time that day.
    assert due_of(task(due_date="2026-09-28T07:00:00Z", bid_closing_date=None)) == datetime(
        2026, 9, 28, 19, 59, 59, tzinfo=UTC
    )
    assert due_of(task(due_date=None, bid_closing_date=None)) is None
    assert not decide(
        task(due_date=None, bid_closing_date=None), now=at(60), grace_minutes=20, watch_from=WATCH
    ).ask


def row(**overrides) -> TaskFollowup:
    fields = dict(
        task_id="901", task_title="test — overdue follow-up", task_url="https://example/DispForm.aspx?ID=901",
        end_user="ADNOC", status_at_ask="Not Started", due_at=DUE,
        assignee_email="sebin@hamdaz.com", status=FollowupStatus.PENDING,
    )
    fields.update(overrides)
    followup = TaskFollowup(**fields)
    followup.assignee = User(email="sebin@hamdaz.com", display_name="Sebin", entra_object_id="x")
    return followup


def test_the_ask_offers_the_way_out_for_a_task_already_updated() -> None:
    link = "https://calm-sky-08cea4100.6.azurestaticapps.net/followups/abc"
    body = mailer.ask_body(row(), link)

    assert "ignore this email" in body
    assert f"{link}?false-positive=1" in body
    assert "Submit Reason" in body and "Already Updated" in body
    # In the Gulf's own time: 10:00 UTC is 14:00 in Dubai.
    assert "14:00 UAE" in body
    assert "past its due time" in body
    assert "AI-generated" in body and "cid:hamdaz-logo" in body
    assert mailer.ask_subject(row()).startswith("Reason Required:")


def test_an_early_ask_says_it_was_marked_not_submitted() -> None:
    body = mailer.ask_body(row(), "https://x/followups/abc", early=True)
    assert "marked <b>Not Submitted</b>" in body
    assert "(Not Submitted)" in mailer.ask_subject(row(), early=True)


def test_mail_links_point_at_production_not_localhost() -> None:
    import uuid as _uuid

    from app.core.config import get_settings
    from app.followups.service import form_link

    link = form_link(get_settings(), _uuid.UUID(int=1))
    assert link.startswith("https://") and "localhost" not in link


def test_the_reason_mail_carries_what_was_said_and_escapes_it() -> None:
    answered = row(reason="Waiting on <supplier> pricing\nsince Friday", answered_at=at(90))
    body = mailer.reason_body(answered, "Sebin", "http://localhost:3000/followups/abc")

    assert "Waiting on &lt;supplier&gt; pricing" in body
    assert "Sebin" in mailer.reason_subject(answered, "Sebin")


def test_today_is_the_gulfs_day_not_utcs() -> None:
    from app.followups.service import today_bounds

    # 22:00 UTC on the 27th is already 02:00 on the 28th in Dubai.
    start, end = today_bounds(datetime(2026, 9, 27, 22, 0, tzinfo=UTC))
    assert start == datetime(2026, 9, 27, 20, 0, tzinfo=UTC)
    assert end == datetime(2026, 9, 28, 20, 0, tzinfo=UTC)


def test_the_screen_says_when_a_task_is_outside_the_trial() -> None:
    from app.followups.service import watches
    from app.models.followup import FollowupSettings

    trial = FollowupSettings(
        enabled=True, only_emails=["sebin@hamdaz.com"], only_title_contains="test",
        watch_from=WATCH,
    )
    sebin = User(email="sebin@hamdaz.com", display_name="Sebin", entra_object_id="x")
    fasna = User(email="fasna@hamdaz.com", display_name="Fasna", entra_object_id="y")

    assert watches(trial, sebin, task()) is None
    assert "outside the trial" in watches(trial, fasna, task())
    assert "in the title" in watches(trial, sebin, task(title="RFQ 6000151129"))
    trial.enabled = False
    assert watches(trial, sebin, task()) == "The follow-up is switched off."


def test_due_today_lists_todays_tasks_soonest_first_with_the_ask_time() -> None:
    from app.followups.service import due_today_rows

    person = User(email="sebin@hamdaz.com", display_name="Sebin", entra_object_id="x")
    later = task(id="902", bid_closing_date=stored(datetime(2026, 9, 28, 15, 0, tzinfo=UTC)))
    done = task(id="903", submission_status="Submitted",
                bid_closing_date=stored(datetime(2026, 9, 28, 6, 0, tzinfo=UTC)))
    tomorrow = task(id="904", bid_closing_date=stored(datetime(2026, 9, 29, 21, 0, tzinfo=UTC)))
    rows = due_today_rows(
        [(person, [later, task(), done, tomorrow])],
        now=datetime(2026, 9, 28, 8, 0, tzinfo=UTC), grace_minutes=20, asked={},
    )

    assert [r["task_id"] for r in rows] == ["901", "902", "903"]
    assert rows[0]["ask_at"] == DUE + timedelta(minutes=20)
    assert rows[-1]["finished"] is True



# ── the end-of-day report ──────────────────────────────────────────────


def _digest_settings(**overrides):
    from app.models.followup import FollowupSettings

    fields = dict(
        digest_enabled=True, digest_time="18:00", digest_timezone="Asia/Kolkata",
        digest_recipients=["sebin@hamdaz.com"], digest_include_ceo=False,
        digest_formats=["pdf", "xlsx"], digest_last_sent_on=None,
        weekly_enabled=True, weekly_day=4, weekly_last_sent_on=None,
    )
    fields.update(overrides)
    return FollowupSettings(**fields)


def test_the_report_goes_at_six_in_the_evening_india_time_once_a_day() -> None:
    from datetime import date

    from app.followups import digest

    row = _digest_settings()
    # 18:00 IST is 12:30 UTC.
    assert digest.cutoff_for(row, date(2026, 9, 29)) == datetime(2026, 9, 29, 12, 30, tzinfo=UTC)
    assert not digest.is_due(row, datetime(2026, 9, 29, 12, 29, tzinfo=UTC))
    assert digest.is_due(row, datetime(2026, 9, 29, 12, 31, tzinfo=UTC))
    row.digest_last_sent_on = date(2026, 9, 29)
    assert not digest.is_due(row, datetime(2026, 9, 29, 15, 0, tzinfo=UTC))
    row.digest_enabled = False
    row.digest_last_sent_on = None
    assert not digest.is_due(row, datetime(2026, 9, 29, 15, 0, tzinfo=UTC))


def _sample_digest():
    from datetime import date

    from app.followups import digest

    built = digest.Digest(
        day=date(2026, 9, 29),
        window_start=datetime(2026, 9, 28, 12, 30, tzinfo=UTC),
        window_end=datetime(2026, 9, 29, 12, 30, tzinfo=UTC),
        zone="Asia/Kolkata", team="Presale",
    )
    built.lines.append(digest.DigestLine(
        person="Rameesa", email="rameesa@hamdaz.com", task="6000151129", end_user="ADNOC",
        due="29 Sep 2026, 19:29", submission_status="Not Submitted", asked="29 Sep 2026, 10:02",
        mailed="Yes", status="Not responded", reason="No reason given by the end of the day.",
        answered="", managers_told="",
    ))
    task = dict(
        email="fasna@hamdaz.com", end_user="ADNOC", due="29 Sep 2026, 17:51", status="Completed",
        current_type="RFP", priority="High", quote_no="QT-000123", order_status="",
        remarks="22/09 - Enquiry sent to OEMs", followup="",
    )
    built.submitted.append(digest.TaskLine(person="Fasna Sherin", task="0020007321 Single Stage RFP",
                                           submission_status="Submitted", **task))
    built.not_submitted.append(digest.TaskLine(person="Fasna Sherin", task="6000150622 Omnis Software",
                                               submission_status="Not Submitted", **task))
    return built


def test_the_report_builds_as_a_pdf_and_a_workbook_with_every_list() -> None:
    import io

    from openpyxl import load_workbook

    from app.followups import digest

    built = _sample_digest()
    pdf = digest.build_pdf(built)
    assert pdf.startswith(b"%PDF") and len(pdf) > 2000

    wb = load_workbook(io.BytesIO(digest.build_xlsx(built)))
    assert wb.sheetnames == ["Not submitted", "Submitted", "Reasons"]
    values = [c.value for row in wb["Not submitted"].iter_rows() for c in row]
    assert "6000150622 Omnis Software" in values and "QT-000123" in values
    assert "Not responded" in [c.value for row in wb["Reasons"].iter_rows() for c in row]


def test_the_mail_leads_with_what_was_not_submitted() -> None:
    from app.followups import digest

    html = digest.mail_html(_sample_digest(), "https://x/followups")
    assert "Not submitted" in html and "6000150622 Omnis Software" in html
    assert "No reason given by the closing time" in html and "Rameesa" in html
    # The body stays short: reasons and remarks are in the attachments.
    assert "Enquiry sent to OEMs" not in html
    assert "View Report" in html and "AI-generated" in html



def test_the_weekly_report_goes_on_friday_at_the_closing_time() -> None:
    from datetime import date

    from app.followups import digest

    row = _digest_settings()
    friday = datetime(2026, 10, 2, 12, 31, tzinfo=UTC)  # 18:01 IST on a Friday
    assert digest.is_weekly_due(row, friday)
    assert not digest.is_weekly_due(row, datetime(2026, 10, 2, 12, 29, tzinfo=UTC))
    assert not digest.is_weekly_due(row, datetime(2026, 10, 1, 12, 31, tzinfo=UTC))  # Thursday
    row.weekly_last_sent_on = date(2026, 10, 2)
    assert not digest.is_weekly_due(row, friday)
    row.weekly_last_sent_on = None
    row.weekly_enabled = False
    assert not digest.is_weekly_due(row, friday)


def test_the_weekly_report_has_a_by_person_summary() -> None:
    import io

    from openpyxl import load_workbook

    from app.followups import digest

    built = _sample_digest()
    built.days = 7
    assert built.title == "Weekly report — 23 – 29 Sep 2026"
    assert built.people()[0][0] == "Fasna Sherin"  # most not submitted first

    wb = load_workbook(io.BytesIO(digest.build_xlsx(built)))
    assert wb.sheetnames == ["By person", "Not submitted", "Submitted", "Reasons"]
    assert digest.build_pdf(built).startswith(b"%PDF")
    html = digest.mail_html(built, "https://x/followups")
    assert "By person" in html and "Due this week" in html
