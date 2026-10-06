"""Enquiry analysis: the parts that need no database, no key and no network."""

from __future__ import annotations

import io
import json
import uuid
from datetime import UTC, date, datetime, timedelta
from decimal import Decimal

import pytest
from openpyxl import load_workbook

from app.comparison.documents import Readable
from app.enquiries import matching, reading, report, research
from app.models.enquiry import EnquiryAnalysis, EnquiryDocument, EnquiryLine

# ── which file is which ────────────────────────────────────────────────


@pytest.mark.parametrize(
    ("name", "kind"),
    [
        ("RFQ 6000150626.pdf", "requirement"),
        ("Request for Quotation - valves.pdf", "requirement"),
        ("BOQ_rev2.xlsx", "requirement"),
        ("Technical Specification.docx", "requirement"),
        ("Quotation QT-4471 router-switch.pdf", "supplier_quote"),
        ("Offer 2231 from Emerson.pdf", "supplier_quote"),
        ("proforma_invoice.pdf", "supplier_quote"),
        ("scan0003.pdf", None),
    ],
)
def test_guess_kind(name: str, kind: str | None) -> None:
    assert reading.guess_kind(name) == kind


def test_own_reports_are_not_read_back() -> None:
    assert reading.is_own_output("QT-001720 Selling & Costing Report pass 1.pdf", "Quote request QT-001720/x.pdf", "Enquiry analysis")
    assert reading.is_own_output("anything.pdf", "Enquiry analysis/anything.pdf", "Enquiry analysis")
    assert not reading.is_own_output("RFQ.pdf", "RFQ.pdf", "Enquiry analysis")


# ── what is sent and what comes back ───────────────────────────────────


def test_schema_is_strict_at_every_level() -> None:
    def walk(node: object) -> None:
        if isinstance(node, dict):
            if node.get("type") == "object":
                assert node.get("additionalProperties") is False
                assert set(node["required"]) == set(node["properties"])
            for value in node.values():
                walk(value)
        elif isinstance(node, list):
            for value in node:
                walk(value)

    walk(reading.SCHEMA)
    walk(research.SCHEMA)


def test_parts_for_sends_text_as_text_and_scans_as_document_blocks() -> None:
    text = Readable("text", "application/pdf", "RFQ.pdf", text="Item 1: valve " * 10)
    scan = Readable("document", "application/pdf", "scan.pdf", data=b"%PDF-1.4 scan")
    photo = Readable("image", "image/png", "photo.png", data=b"\x89PNG")
    out = reading.parts_for([(text, "task attachment"), (scan, "task folder: scan.pdf"), (photo, "upload")], context="Enquiry: X")
    types = [p["type"] for p in out.parts]
    assert types[0] == "text" and "Enquiry: X" in out.parts[0]["text"]
    assert "document" in types and "image" in types
    assert out.names == ["RFQ.pdf", "scan.pdf", "photo.png"]
    document = out.parts[types.index("document")]
    assert document["source"]["type"] == "base64" and document["source"]["media_type"] == "application/pdf"
    assert document["title"] == "scan.pdf"


def test_trim_keeps_head_and_tail() -> None:
    text = "A" * 1000 + "B" * 1000
    cut = reading.trim(text, 500)
    assert cut.startswith("A") and cut.endswith("B") and "left out" in cut


def test_parse_drops_blank_and_duplicate_items() -> None:
    raw = json.dumps(
        {
            "summary": "Valves for a plant.",
            "customer": "ADNOC",
            "deadline": "2026-10-20",
            "conditions": ["ISO 9001", " "],
            "missing": ["Delivery address"],
            "documents": [{"file_name": "RFQ.pdf", "kind": "requirement"}, {"file_name": "Q.pdf", "kind": "supplier_quote"}],
            "items": [
                {"description": "Ball valve 2in", "part_number": "BV-2", "brand": "", "quantity": 4, "unit": "nos",
                 "specification": "", "source_document": "RFQ.pdf"},
                {"description": "Ball valve 2in", "part_number": "BV-2", "brand": "", "quantity": 4, "unit": "nos",
                 "specification": "", "source_document": "BOQ.xlsx"},
                {"description": "  ", "part_number": "", "brand": "", "quantity": 1, "unit": "",
                 "specification": "", "source_document": ""},
                {"description": "Gasket", "part_number": "", "brand": "Klinger", "quantity": 0, "unit": "",
                 "specification": "", "source_document": "RFQ.pdf"},
            ],
        }
    )
    out = reading.parse(raw)
    assert [i.description for i in out.items] == ["Ball valve 2in", "Gasket"]
    assert out.items[0].quantity == Decimal("4.0000")
    assert out.items[1].quantity is None and out.items[1].part_number is None
    assert out.conditions == ["ISO 9001"]
    assert out.kinds == {"RFQ.pdf": "requirement", "Q.pdf": "supplier_quote"}


def test_parse_refuses_prose() -> None:
    with pytest.raises(ValueError):
        reading.parse("Here are the items: ...")


# ── matching ───────────────────────────────────────────────────────────


def test_part_numbers_match_without_punctuation() -> None:
    assert matching.norm_pn("C9300-48P-E") == matching.norm_pn("c9300 48p e")
    words = matching.tokens("Catalyst switch")
    assert matching.score("C930048PE", words, None, "C930048PE", frozenset(), None) == 1.0


def test_descriptions_match_by_covering_the_requirement() -> None:
    want = matching.tokens("Ball valve 2 inch stainless steel 316 flanged")
    same = matching.tokens("2 inch ball valve, SS 316, flanged ends, stainless steel, ANSI 150")
    other = matching.tokens("Gate valve 6 inch carbon steel")
    assert matching.score("", want, None, "", same, None) >= matching.THRESHOLD
    assert matching.score("", want, None, "", other, None) < matching.THRESHOLD


def test_a_different_brand_counts_against() -> None:
    want = matching.tokens("pressure transmitter 0-10 bar 4-20mA")
    have = matching.tokens("pressure transmitter 0-10 bar 4-20mA HART")
    same_brand = matching.score("", want, "Rosemount", "", have, "Rosemount")
    other_brand = matching.score("", want, "Rosemount", "", have, "Yokogawa")
    assert other_brand < same_brand


def _known(**kw) -> matching.Known:
    base = dict(
        source="supplier_quote", ref="Q1", when=date(2026, 9, 1), counterparty="Acme",
        description="Ball valve 2 inch SS316 flanged", part_number=None, brand=None,
        rate=Decimal("120"), currency="AED", supplier="Acme",
    )
    base.update(kw)
    return matching.Known(**base)


def test_pool_finds_by_part_number_and_words_and_orders_best_first() -> None:
    pool = matching.Pool(
        [
            _known(ref="old", when=date(2023, 1, 1)),
            _known(ref="pn", part_number="BV-316-2F", description="Valve"),
            _known(ref="unrelated", description="Cable tray 300mm galvanised"),
        ]
    )
    found = pool.match("Ball valve 2 inch stainless SS316 flanged", "bv 316 2f", None)
    refs = [row.ref for _score, row in found]
    assert refs[0] == "pn"
    assert "old" in refs and "unrelated" not in refs


def test_status_by_latest_date() -> None:
    today = date(2026, 10, 5)
    recent = [{"date": (today - timedelta(days=40)).isoformat()}]
    old = [{"date": "2023-02-01"}, {"date": None}]
    assert matching.status_of([], recent_months=12, today=today) == "new"
    assert matching.status_of(recent, recent_months=12, today=today) == "recent"
    assert matching.status_of(old, recent_months=12, today=today) == "history"


def test_suppliers_one_per_name_latest_first() -> None:
    history = [
        {"source": "supplier_quote", "supplier": "Acme LLC", "rate": "100", "currency": "AED", "date": "2025-01-01"},
        {"source": "supplier_quote", "supplier": "acme llc", "rate": "110", "currency": "AED", "date": "2026-01-01"},
        {"source": "zoho_po", "counterparty": "Beta Trading", "currency": "USD", "date": "2024-06-01"},
        {"source": "zoho_estimate", "counterparty": "A Customer", "date": "2026-02-01"},
    ]
    current = [(0.9, _known(supplier="This Enquiry Supplier", when=date(2026, 10, 1), current=True))]
    out = matching.suppliers_from(history, current)
    names = [s["name"] for s in out]
    assert names == ["This Enquiry Supplier", "acme llc", "Beta Trading"]
    assert out[1]["last_rate"] == "110"
    assert all(s["partner"] is None for s in out)


# ── the web lookup's answer ────────────────────────────────────────────


def test_research_parse() -> None:
    raw = json.dumps(
        {
            "items": [
                {
                    "index": 1, "manufacturer": "Emerson", "manufacturer_website": "https://emerson.com",
                    "product_url": "", "price_low": 250, "price_high": 0, "currency": "usd",
                    "price_source": "distributor page", "notes": "", "confidence": 1.4,
                    "suppliers": [
                        {"name": "Gulf Valves", "role": "distributor", "website": "", "email": "", "phone": "",
                         "country": "UAE", "evidence": "lists the model"},
                        {"name": " ", "role": "unknown", "website": "", "email": "", "phone": "", "country": "",
                         "evidence": ""},
                    ],
                },
                {"index": "x"},
            ]
        }
    )
    out = research.parse(raw)
    assert list(out) == [1]
    assert out[1]["price_low"] == "250.00" and out[1]["price_high"] == "250.00"
    assert out[1]["currency"] == "USD" and out[1]["confidence"] == 1.0
    assert [s["name"] for s in out[1]["suppliers"]] == ["Gulf Valves"]
    assert out[1]["suppliers"][0]["source"] == "web" and out[1]["suppliers"][0]["email"] is None
    assert research.parse("not json") == {}


def test_no_price_means_no_currency() -> None:
    raw = json.dumps({"items": [{"index": 1, "manufacturer": "", "manufacturer_website": "", "product_url": "",
                                 "price_low": 0, "price_high": 0, "currency": "AED", "price_source": "",
                                 "suppliers": [], "notes": "", "confidence": 0.2}]})
    assert research.parse(raw)[1]["currency"] is None


# ── the reports ────────────────────────────────────────────────────────


def _analysis() -> EnquiryAnalysis:
    analysis = EnquiryAnalysis(
        id=uuid.uuid4(), task_id="123", task_title="6000150626 Valves for <Ruwais> & co",
        end_user="ADNOC", bid_closing_date=date(2026, 10, 20), status="done",
        summary="Valves and gaskets.", conditions=["ISO 9001"], missing=["Delivery address"],
        run_notes=["1 new item was not looked up"], finished_at=datetime(2026, 10, 5, tzinfo=UTC),
        documents=[], lines=[],
    )
    analysis.documents = [
        EnquiryDocument(source="attachment", origin_key="attachment:RFQ.pdf", file_name="RFQ.pdf",
                        kind="requirement", status="read"),
        EnquiryDocument(source="folder", origin_key="drive:1", file_name="Quotation.pdf", path="Quotation.pdf",
                        kind="supplier_quote", status="read", note="Stored", supplier_quote_id=uuid.uuid4()),
    ]
    analysis.lines = [
        EnquiryLine(
            position=0, description="Ball valve 2in", part_number="BV-2", brand="Emerson",
            quantity=Decimal("4"), unit="nos", status="recent",
            history=[{"source": "supplier_quote", "ref": "Q1", "date": "2026-09-01", "supplier": "Acme",
                      "rate": "120", "currency": "AED", "score": 1.0}],
            suppliers=[{"name": "Acme", "source": "supplier_quote", "last_rate": "120", "currency": "AED",
                        "last_date": "2026-09-01", "partner": None}],
            web=None,
        ),
        EnquiryLine(
            position=1, description="Spiral wound gasket", quantity=None, status="new", history=[],
            suppliers=[{"name": "Gulf Seals", "source": "web", "website": "https://x", "partner": None}],
            web={"manufacturer": "Klinger", "price_low": "10.00", "price_high": "14.00", "currency": "USD"},
        ),
    ]
    return analysis


def test_counts_and_latest() -> None:
    a = _analysis()
    assert report.counts(a) == {"recent": 1, "history": 0, "new": 1}
    assert report.latest(a.lines[0])["ref"] == "Q1"
    assert report.latest(a.lines[1]) is None


def test_pdf_renders_with_awkward_characters() -> None:
    content = report.pdf(_analysis())
    assert content.startswith(b"%PDF")


def test_workbook_has_a_sheet_per_view() -> None:
    book = load_workbook(io.BytesIO(report.workbook(_analysis())))
    assert book.sheetnames == ["Summary", "Items", "History", "Suppliers", "Documents"]
    items = book["Items"]
    assert items.max_row == 3
    assert items["G2"].value == "Seen recently" and items["G3"].value == "New item"
    assert book["Suppliers"]["I2"].value == "not recorded"


def test_file_stem_is_short() -> None:
    a = _analysis()
    a.task_title = "x" * 300
    assert len(report.file_stem(a)) < 120


# ── listing a task folder (stubbed Graph; nothing live) ────────────────


async def test_list_files_walks_subfolders_and_pages() -> None:
    import httpx

    from app.core.config import get_settings
    from app.quoting.storage import QuoteDrive

    def handler(request: httpx.Request) -> httpx.Response:
        url = str(request.url)
        if url.endswith("/oauth2/v2.0/token"):
            return httpx.Response(200, json={"access_token": "t", "expires_in": 3600})
        path = httpx.URL(url).path
        if "page2" in url:
            return httpx.Response(200, json={"value": [{"id": "f3", "name": "BOQ.xlsx", "file": {}, "size": 9}]})
        if path.endswith("Task 1:/children"):
            return httpx.Response(
                200,
                json={
                    "value": [
                        {"id": "f1", "name": "RFQ.pdf", "file": {"mimeType": "application/pdf"}, "size": 10},
                        {"id": "d1", "name": "Quote request QT-1", "folder": {}},
                    ],
                    "@odata.nextLink": "https://graph.microsoft.com/v1.0/next?page2",
                },
            )
        if path.endswith("Task 1/Quote request QT-1:/children"):
            return httpx.Response(200, json={"value": [{"id": "f2", "name": "Quotation.pdf", "file": {}, "size": 5}]})
        return httpx.Response(404, json={})

    settings = get_settings().model_copy()
    settings.quote_drive_id = "drive"
    settings.quote_drive_folder = "Proposal Team Channel"
    drive = QuoteDrive(settings, httpx.AsyncClient(transport=httpx.MockTransport(handler)))
    files = await drive.list_files("Task 1")
    assert [(f["id"], f["path"]) for f in files] == [
        ("f1", "RFQ.pdf"),
        ("f3", "BOQ.xlsx"),
        ("f2", "Quote request QT-1/Quotation.pdf"),
    ]
    assert await drive.list_files("Missing") == []


# ── the Claude client, without a network ──────────────────────────────


def test_claude_cost_and_errors() -> None:
    import anthropic
    import httpx

    from app.enquiries import claude

    assert claude.cost_of("claude-opus-5-5", 1_000_000, 100_000) == Decimal("6.0000")
    assert claude.cost_of("unknown-model", 1000, 1000) == Decimal("0.0000")
    request = httpx.Request("POST", "https://api.anthropic.com/v1/messages")
    bad_key = anthropic.AuthenticationError("invalid", response=httpx.Response(401, request=request), body=None)
    assert "ANTHROPIC_API_KEY" in claude.explain(bad_key)
    broke = anthropic.BadRequestError(
        "Your credit balance is too low to access the Anthropic API.",
        response=httpx.Response(400, request=request), body=None,
    )
    assert "no credit" in claude.explain(broke)


async def test_research_resumes_a_paused_search_and_reads_the_tool() -> None:
    from types import SimpleNamespace

    from app.core.config import get_settings
    from app.enquiries.claude import ClaudeEnquiry

    usage = SimpleNamespace(input_tokens=10, output_tokens=5, cache_creation_input_tokens=0, cache_read_input_tokens=0)
    replies = [
        SimpleNamespace(stop_reason="pause_turn", usage=usage, content=[SimpleNamespace(type="server_tool_use", name="web_search")]),
        SimpleNamespace(stop_reason="tool_use", usage=usage,
           content=[SimpleNamespace(type="tool_use", name="record_findings", input={"items": []})]),
    ]
    sent: list[list] = []

    class FakeMessages:
        async def create(self, **kw):
            sent.append([m["role"] for m in kw["messages"]])
            assert kw["fallbacks"] == "default" and kw["tools"][0]["type"] == "web_search_20260209"
            return replies.pop(0)

    settings = get_settings().model_copy()
    settings.anthropic_api_key = "test"
    llm = ClaudeEnquiry(settings)
    llm._client = SimpleNamespace(beta=SimpleNamespace(messages=FakeMessages()))
    raw, tokens_in, tokens_out = await llm.research(instructions="x", prompt="1. valve", schema={}, user_key="u")
    assert json.loads(raw) == {"items": []}
    assert (tokens_in, tokens_out) == (20, 10)
    assert sent == [["user"], ["user", "assistant"]]
