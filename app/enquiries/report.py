"""The enquiry analysis as a PDF and as a workbook.

The PDF is for reading: what the customer wants, which items we have met and
where, who might supply the rest, and what the documents leave unsaid. The
workbook is for working: one row per line, then every history entry and every
candidate supplier on sheets of their own, ready to filter and price.

Both carry the same warning wherever a web figure appears — a price found on a
public page is a guide, not a quotation.
"""

from __future__ import annotations

import io
import re
from datetime import UTC, datetime
from decimal import Decimal
from typing import Any, Final

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.utils import ImageReader
from reportlab.platypus import KeepTogether, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from app.models.enquiry import EnquiryAnalysis, EnquiryLine
from app.quoting.report_pdf import (
    ADDRESS,
    CYAN,
    FAINT,
    FOOTER,
    INK,
    LINE,
    LOGO,
    MAGENTA,
    MUTED,
    NAVY,
    ORG,
    ROW,
    ROW_BLUE,
)

STATUS_LABEL: Final = {
    "recent": "Seen recently",
    "history": "In our history",
    "new": "New item",
}
SOURCE_LABEL: Final = {
    "supplier_quote": "Supplier quote",
    "quote_request": "Our quote",
    "zoho_item": "Zoho item",
    "zoho_po": "Zoho PO",
    "zoho_bill": "Zoho bill",
    "zoho_estimate": "Zoho quote",
    "this_enquiry": "Offered for this enquiry",
    "web": "Web",
}
WEB_WARNING: Final = "Web prices are from public pages: a guide, not a quotation."

GREEN = colors.HexColor("#e5f5ec")
GREEN_TEXT = colors.HexColor("#17663a")
AMBER = colors.HexColor("#fdf1dc")
AMBER_TEXT = colors.HexColor("#8a5a00")
BLUE = colors.HexColor("#e6f3fa")


def file_stem(analysis: EnquiryAnalysis) -> str:
    """"<task title> enquiry analysis", short enough for a library path."""
    title = re.sub(r"\s+", " ", analysis.task_title or "Enquiry").strip()[:90].rstrip()
    return f"{title} - enquiry analysis"


def _when(analysis: EnquiryAnalysis) -> str:
    stamp = analysis.finished_at or datetime.now(UTC)
    return stamp.strftime("%d %b %Y")


def _num(value: Any) -> str:
    if value in (None, ""):
        return "—"
    try:
        number = Decimal(str(value))
    except Exception:  # noqa: BLE001
        return str(value)
    if number == number.to_integral_value():
        return f"{number:,.0f}"
    return f"{number:,.2f}"


def latest(line: EnquiryLine) -> dict[str, Any] | None:
    """The most recent history entry, which is what "last seen" means."""
    entries = [h for h in line.history or [] if h.get("date")]
    if not entries:
        return (line.history or [None])[0]
    return max(entries, key=lambda h: h["date"])


def counts(analysis: EnquiryAnalysis) -> dict[str, int]:
    out = {"recent": 0, "history": 0, "new": 0}
    for line in analysis.lines:
        out[line.status] = out.get(line.status, 0) + 1
    return out


# ── the PDF ────────────────────────────────────────────────────────────

PAGE_W, PAGE_H = landscape(A4)
MARGIN: Final = 28.0
WIDTH: Final = PAGE_W - 2 * MARGIN


def _style(size: float = 7.6, *, bold: bool = False, color: Any = INK, leading: float | None = None) -> ParagraphStyle:
    return ParagraphStyle(
        "s",
        fontName="Helvetica-Bold" if bold else "Helvetica",
        fontSize=size,
        leading=leading or size * 1.28,
        textColor=color,
    )


def _p(text: Any, size: float = 7.6, **kw: Any) -> Paragraph:
    raw = "" if text is None else str(text)
    safe = raw.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;").replace("\n", "<br/>")
    return Paragraph(safe, _style(size, **kw))


def _page(analysis: EnquiryAnalysis):
    logo = ImageReader(str(LOGO)) if LOGO.exists() else None
    footer = ImageReader(str(FOOTER)) if FOOTER.exists() else None

    def draw(canvas, doc) -> None:
        canvas.saveState()
        y = PAGE_H
        if logo is not None:
            canvas.drawImage(logo, MARGIN, y - 72, width=56, height=44, mask="auto")
        canvas.setFont("Helvetica-Bold", 9)
        canvas.setFillColor(NAVY)
        canvas.drawString(MARGIN + 66, y - 36, ORG)
        canvas.setFont("Helvetica", 7.4)
        canvas.setFillColor(MUTED)
        for i, line in enumerate(ADDRESS[:3]):
            canvas.drawString(MARGIN + 66, y - 49 - i * 10.5, line)

        canvas.setFont("Helvetica-Bold", 14)
        canvas.setFillColor(NAVY)
        canvas.drawRightString(PAGE_W - MARGIN, y - 40, "ENQUIRY ANALYSIS")
        rule = 160.0
        canvas.setFillColor(CYAN)
        canvas.rect(PAGE_W - MARGIN - rule, y - 48, rule * 0.62, 2.2, stroke=0, fill=1)
        canvas.setFillColor(MAGENTA)
        canvas.rect(PAGE_W - MARGIN - rule * 0.38, y - 48, rule * 0.38, 2.2, stroke=0, fill=1)
        canvas.setFont("Helvetica", 8)
        canvas.setFillColor(MUTED)
        canvas.drawRightString(PAGE_W - MARGIN, y - 62, f"Prepared {_when(analysis)}")
        canvas.setStrokeColor(LINE)
        canvas.setLineWidth(0.8)
        canvas.line(MARGIN, y - 82, PAGE_W - MARGIN, y - 82)

        base = 14.0
        if footer is not None:
            iw, ih = footer.getSize()
            fh = min(36.0, WIDTH * ih / iw)
            fw = fh * iw / ih
            canvas.drawImage(footer, MARGIN, 12, width=fw, height=fh, mask="auto")
            base = 12 + fh + 4
        canvas.setFont("Helvetica", 6.5)
        canvas.setFillColor(FAINT)
        canvas.drawString(MARGIN, base, f"Enquiry analysis · {analysis.task_title[:110]}")
        canvas.drawRightString(PAGE_W - MARGIN, base, f"Page {doc.page}")
        canvas.restoreState()

    return draw


def _section(text: str) -> Paragraph:
    return Paragraph(text.upper(), _style(8.4, bold=True, color=NAVY, leading=12))


def _facts(analysis: EnquiryAnalysis) -> Table:
    c = counts(analysis)
    quotes = sum(1 for d in analysis.documents if d.kind == "supplier_quote" and d.supplier_quote_id)
    rows = [
        [_p("Enquiry", 6.8, color=MUTED), _p(analysis.task_title, 9, bold=True)],
        [_p("Customer / end user", 6.8, color=MUTED), _p(analysis.customer or analysis.end_user or "—", 8.4)],
        [_p("Bid closing", 6.8, color=MUTED), _p(
            analysis.bid_closing_date.strftime("%d %b %Y") if analysis.bid_closing_date else (analysis.deadline or "—"), 8.4
        )],
    ]
    left = Table(rows, colWidths=[90, WIDTH * 0.55 - 90])
    left.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "TOP"), ("BOTTOMPADDING", (0, 0), (-1, -1), 3)]))

    def tile(label: str, value: int, fill: Any, ink: Any) -> Table:
        t = Table([[_p(str(value), 15, bold=True, color=ink)], [_p(label, 6.8, color=ink)]], colWidths=[WIDTH * 0.45 / 4 - 6])
        t.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, -1), fill), ("LEFTPADDING", (0, 0), (-1, -1), 7)]))
        return t

    tiles = Table(
        [[
            tile("Items asked for", len(analysis.lines), ROW_BLUE, NAVY),
            tile("Seen recently", c["recent"], GREEN, GREEN_TEXT),
            tile("In our history", c["history"], BLUE, NAVY),
            tile("New items", c["new"], AMBER, AMBER_TEXT),
        ]],
        colWidths=[WIDTH * 0.45 / 4] * 4,
    )
    outer = Table([[left, tiles]], colWidths=[WIDTH * 0.55, WIDTH * 0.45])
    outer.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "TOP"), ("LEFTPADDING", (0, 0), (-1, -1), 0)]))
    if quotes:
        return Table([[outer], [_p(f"{quotes} supplier quotation(s) found among the documents and stored.", 7.2, color=MUTED)]])
    return outer


def _bullets(items: list[str] | None) -> list[Paragraph]:
    return [_p(f"•  {i}", 7.8) for i in items or []]


def _seen(line: EnquiryLine) -> str:
    h = latest(line)
    if not h:
        return "—"
    parts = [SOURCE_LABEL.get(h.get("source"), h.get("source") or ""), h.get("ref") or ""]
    if h.get("date"):
        parts.append(str(h["date"]))
    who = h.get("supplier") or h.get("counterparty")
    if who:
        parts.append(who)
    rate = h.get("rate") or h.get("cost_rate")
    if rate:
        parts.append(f"{h.get('currency') or ''} {_num(rate)}".strip())
    more = len(line.history or []) - 1
    text = " · ".join(p for p in parts if p)
    return f"{text}\n+{more} more" if more > 0 else text


def _suppliers(line: EnquiryLine, limit: int = 4) -> str:
    out = []
    for s in (line.suppliers or [])[:limit]:
        bits = [s.get("name") or ""]
        if s.get("last_rate"):
            bits.append(f"{s.get('currency') or ''} {_num(s['last_rate'])}".strip())
        source = SOURCE_LABEL.get(s.get("source"), s.get("source") or "")
        if source:
            bits.append(f"({source})")
        if s.get("partner") is True:
            bits.append("· partner")
        out.append(" ".join(b for b in bits if b))
    more = len(line.suppliers or []) - limit
    if more > 0:
        out.append(f"+{more} more")
    return "\n".join(out) or "—"


def _web(line: EnquiryLine) -> str:
    web = line.web or {}
    if not web:
        return "—"
    parts = []
    if web.get("manufacturer"):
        parts.append(f"Maker: {web['manufacturer']}")
    low, high = web.get("price_low"), web.get("price_high")
    if low:
        rng = _num(low) if not high or high == low else f"{_num(low)}–{_num(high)}"
        parts.append(f"Web price: {web.get('currency') or ''} {rng}".strip())
    if web.get("notes"):
        parts.append(str(web["notes"])[:160])
    return "\n".join(parts) or "Nothing found"


def _lines_table(analysis: EnquiryAnalysis) -> Table:
    widths = [18, 196, 44, 62, 160, 176, WIDTH - 18 - 196 - 44 - 62 - 160 - 176]
    head = ["#", "Item", "Qty", "Status", "Last seen", "Suppliers", "Web lookup"]
    rows: list[list[Any]] = [[_p(h, 6.8, bold=True, color=colors.white) for h in head]]
    fills = []
    for n, line in enumerate(analysis.lines, start=1):
        item = line.description
        extra = " · ".join(x for x in (line.part_number and f"P/N {line.part_number}", line.brand) if x)
        if extra:
            item += f"\n{extra}"
        status = STATUS_LABEL.get(line.status, line.status)
        rows.append([
            _p(n, 7), _p(item, 7.2), _p(f"{_num(line.quantity)} {line.unit or ''}".strip(), 7),
            _p(status, 7, bold=True, color=GREEN_TEXT if line.status == "recent" else AMBER_TEXT if line.status == "new" else NAVY),
            _p(_seen(line), 6.8), _p(_suppliers(line), 6.8), _p(_web(line), 6.8),
        ])
        if n % 2 == 0:
            fills.append(("BACKGROUND", (0, n), (-1, n), ROW))
    table = Table(rows, colWidths=widths, repeatRows=1)
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), NAVY),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("LINEBELOW", (0, 1), (-1, -1), 0.4, LINE),
        ("TOPPADDING", (0, 0), (-1, -1), 3),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
        *fills,
    ]))
    return table


def _documents_table(analysis: EnquiryAnalysis) -> Table:
    rows: list[list[Any]] = [[_p(h, 6.8, bold=True, color=colors.white) for h in ("Document", "From", "Read as", "Note")]]
    for d in analysis.documents:
        rows.append([
            _p(d.path or d.file_name, 7), _p({"attachment": "List attachment", "folder": "Task folder", "upload": "Uploaded"}.get(d.source, d.source), 7),
            _p({"requirement": "Requirement", "supplier_quote": "Supplier quote", "other": "Other"}.get(d.kind or "", d.status), 7),
            _p(d.note or "", 6.8, color=MUTED),
        ])
    table = Table(rows, colWidths=[WIDTH * 0.38, WIDTH * 0.12, WIDTH * 0.12, WIDTH * 0.38], repeatRows=1)
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), NAVY),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("LINEBELOW", (0, 1), (-1, -1), 0.4, LINE),
    ]))
    return table


def pdf(analysis: EnquiryAnalysis) -> bytes:
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer, pagesize=landscape(A4), leftMargin=MARGIN, rightMargin=MARGIN,
        topMargin=92, bottomMargin=58, title=file_stem(analysis), author=ORG,
    )
    story: list[Any] = [_facts(analysis), Spacer(1, 10)]
    if analysis.summary:
        story += [_section("Summary"), Spacer(1, 3), _p(analysis.summary, 8.2), Spacer(1, 8)]
    if analysis.lines:
        story += [_section("Items asked for"), Spacer(1, 4), _lines_table(analysis),
                  Spacer(1, 3), _p(WEB_WARNING, 6.6, color=FAINT), Spacer(1, 10)]
    else:
        story += [_p("No items were found in the documents.", 8.2, color=MUTED), Spacer(1, 8)]
    if analysis.conditions:
        story += [KeepTogether([_section("Conditions"), Spacer(1, 3), *_bullets(analysis.conditions)]), Spacer(1, 8)]
    if analysis.missing:
        story += [KeepTogether([_section("Not said in the documents"), Spacer(1, 3), *_bullets(analysis.missing)]), Spacer(1, 8)]
    if analysis.documents:
        story += [_section("Documents read"), Spacer(1, 4), _documents_table(analysis), Spacer(1, 8)]
    if analysis.run_notes:
        story += [KeepTogether([_section("Notes on this run"), Spacer(1, 3), *_bullets(analysis.run_notes)])]
    draw = _page(analysis)
    doc.build(story, onFirstPage=draw, onLaterPages=draw)
    return buffer.getvalue()


# ── the workbook ───────────────────────────────────────────────────────

_HEAD_FILL = PatternFill("solid", fgColor="0E5E80")
_HEAD_FONT = Font(bold=True, color="FFFFFF")
_WRAP = Alignment(wrap_text=True, vertical="top")


def _sheet(book: Workbook, title: str, headers: list[str], rows: list[list[Any]], widths: list[int]) -> None:
    ws = book.create_sheet(title)
    ws.append(headers)
    for cell in ws[1]:
        cell.fill, cell.font, cell.alignment = _HEAD_FILL, _HEAD_FONT, _WRAP
    for row in rows:
        ws.append(row)
    for i, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = _WRAP
    ws.freeze_panes = "A2"
    if rows:
        ws.auto_filter.ref = ws.dimensions


def _dec(value: Any) -> float | None:
    if value in (None, ""):
        return None
    try:
        return float(Decimal(str(value)))
    except Exception:  # noqa: BLE001
        return None


def workbook(analysis: EnquiryAnalysis) -> bytes:
    book = Workbook()
    summary = book.active
    summary.title = "Summary"
    c = counts(analysis)
    for label, value in (
        ("Enquiry", analysis.task_title),
        ("Customer / end user", analysis.customer or analysis.end_user or ""),
        ("Bid closing", analysis.bid_closing_date.isoformat() if analysis.bid_closing_date else analysis.deadline or ""),
        ("Prepared", _when(analysis)),
        ("Items asked for", len(analysis.lines)),
        ("Seen recently", c["recent"]),
        ("In our history", c["history"]),
        ("New items", c["new"]),
        ("Summary", analysis.summary or ""),
        ("Conditions", "\n".join(analysis.conditions or [])),
        ("Not said in the documents", "\n".join(analysis.missing or [])),
        ("Notes on this run", "\n".join(analysis.run_notes or [])),
        ("Note", WEB_WARNING),
    ):
        summary.append([label, value])
    summary.column_dimensions["A"].width = 26
    summary.column_dimensions["B"].width = 110
    for row in summary.iter_rows():
        row[0].font = Font(bold=True, color="0E5E80")
        for cell in row:
            cell.alignment = _WRAP

    line_rows, history_rows, supplier_rows = [], [], []
    for n, line in enumerate(analysis.lines, start=1):
        h = latest(line) or {}
        web = line.web or {}
        line_rows.append([
            n, line.description, line.part_number, line.brand, _dec(line.quantity), line.unit,
            STATUS_LABEL.get(line.status, line.status),
            SOURCE_LABEL.get(h.get("source"), h.get("source")), h.get("ref"), h.get("date"),
            h.get("supplier") or h.get("counterparty"), _dec(h.get("rate")), _dec(h.get("cost_rate")), h.get("currency"),
            ", ".join(s.get("name") or "" for s in line.suppliers or []),
            web.get("manufacturer"), _dec(web.get("price_low")), _dec(web.get("price_high")), web.get("currency"),
            web.get("product_url"), line.specification, line.source_document,
        ])
        for entry in line.history or []:
            history_rows.append([
                n, line.description, SOURCE_LABEL.get(entry.get("source"), entry.get("source")), entry.get("ref"),
                entry.get("date"), entry.get("counterparty"), entry.get("supplier"), entry.get("description"),
                entry.get("part_number"), _dec(entry.get("quantity")), _dec(entry.get("rate")),
                _dec(entry.get("cost_rate")), entry.get("currency"), entry.get("score"),
            ])
        for s in line.suppliers or []:
            supplier_rows.append([
                n, line.description, s.get("name"), SOURCE_LABEL.get(s.get("source"), s.get("source")),
                s.get("role"), _dec(s.get("last_rate")), s.get("currency"), s.get("last_date"),
                "yes" if s.get("partner") is True else "not recorded" if s.get("partner") is None else "no",
                s.get("website"), s.get("email"), s.get("phone"), s.get("country"), s.get("evidence"),
            ])

    _sheet(
        book, "Items",
        ["#", "Description", "Part number", "Brand", "Qty", "Unit", "Status", "Last seen in", "Reference",
         "Date", "Supplier / customer", "Rate", "Cost rate", "Currency", "Suppliers", "Web: maker",
         "Web: price low", "Web: price high", "Web: currency", "Web: product page", "Specification", "From document"],
        line_rows, [5, 48, 18, 14, 8, 8, 14, 14, 16, 11, 24, 11, 11, 9, 36, 20, 11, 11, 9, 30, 48, 26],
    )
    _sheet(
        book, "History",
        ["#", "Item asked for", "Where", "Reference", "Date", "Supplier / customer", "Supplier", "Description there",
         "Part number there", "Qty", "Rate", "Cost rate", "Currency", "Match"],
        history_rows, [5, 40, 14, 16, 11, 24, 24, 44, 18, 8, 11, 11, 9, 8],
    )
    _sheet(
        book, "Suppliers",
        ["#", "Item asked for", "Supplier", "Found in", "Role", "Last rate", "Currency", "Last date", "Partner",
         "Website", "Email", "Phone", "Country", "Evidence"],
        supplier_rows, [5, 40, 30, 18, 13, 11, 9, 11, 12, 28, 28, 18, 14, 44],
    )
    _sheet(
        book, "Documents",
        ["Document", "From", "Read as", "Status", "Note", "Link"],
        [[d.path or d.file_name, d.source, d.kind, d.status, d.note, d.web_url] for d in analysis.documents],
        [50, 12, 16, 10, 60, 50],
    )
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()
