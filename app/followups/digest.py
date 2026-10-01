"""The end-of-day report: every reason asked for today, in one mail to the CEO.

The managers hear each reason as it arrives. The CEO should not — one mail per
late bid is a mailbox nobody reads. So once a day, at the closing time a super
admin sets (6:00 PM India time to begin with), this gathers the day's
follow-ups into one report, marks anybody who has still not answered as **not
responded**, and mails it as a PDF and an Excel workbook.

**Not responded is a mark, not a closure.** The person can still answer
afterwards, and the banner in the app keeps asking them to: the reason is
mandatory, and the end of the day only decides what the CEO is told about it.

**The house template.** The PDF wears the selling & costing report's own
letterhead — the logo, the navy, the cyan and magenta rule, the address block,
the certification strip — drawn from the same constants and assets, so the two
documents cannot drift apart. The workbook uses the same navy for its header.

**The window** is the day as the settings' timezone has it: from yesterday's
closing time to today's. A follow-up asked after today's closing belongs to
tomorrow's report.
"""

from __future__ import annotations

import io
import logging
from dataclasses import dataclass, field
from datetime import UTC, date, datetime, time, timedelta
from html import escape
from typing import Any, Final
from zoneinfo import ZoneInfo

from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.mail import Attachment
from app.models.followup import FollowupSettings, FollowupStatus, TaskFollowup
from app.models.role import Role, UserRole
from app.models.user import User

logger = logging.getLogger("hamdaz.followups.digest")

PDF_TYPE: Final = "application/pdf"
XLSX_TYPE: Final = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
_GULF: Final = ZoneInfo("Asia/Dubai")

#: How each state reads in the report.
STATUS_LABELS: Final = {
    FollowupStatus.PENDING: "Waiting",
    FollowupStatus.NO_RESPONSE: "Not responded",
    FollowupStatus.ANSWERED: "Reason given",
    FollowupStatus.FALSE_POSITIVE: "False positive",
    FollowupStatus.RESOLVED: "Closed",
}


def _zone(name: str | None) -> ZoneInfo:
    try:
        return ZoneInfo(name or "Asia/Kolkata")
    except Exception:  # noqa: BLE001 - a bad name falls back rather than failing the day
        return ZoneInfo("Asia/Kolkata")


def closing_time(row: FollowupSettings) -> time:
    try:
        hour, minute = (int(x) for x in (row.digest_time or "18:00").split(":")[:2])
        return time(hour, minute)
    except ValueError:
        return time(18, 0)


def cutoff_for(row: FollowupSettings, day: date) -> datetime:
    """``day``'s closing time, as a UTC instant."""
    local = datetime.combine(day, closing_time(row), tzinfo=_zone(row.digest_timezone))
    return local.astimezone(UTC)


def local_today(row: FollowupSettings, now: datetime) -> date:
    return now.astimezone(_zone(row.digest_timezone)).date()


WEEKDAYS: Final = ("Monday", "Tuesday", "Wednesday", "Thursday", "Friday", "Saturday", "Sunday")


def is_weekly_due(row: FollowupSettings, now: datetime) -> bool:
    """Whether this week's report should go now: the chosen weekday, past the
    closing time, not yet sent today."""
    if not row.weekly_enabled:
        return False
    today = local_today(row, now)
    if today.weekday() != (row.weekly_day if row.weekly_day is not None else 4):
        return False
    if row.weekly_last_sent_on == today:
        return False
    return now >= cutoff_for(row, today)


def is_due(row: FollowupSettings, now: datetime) -> bool:
    """Whether today's report should go now: past the closing time, not yet sent."""
    if not row.digest_enabled:
        return False
    today = local_today(row, now)
    if row.digest_last_sent_on == today:
        return False
    return now >= cutoff_for(row, today)


#: Under the carried-over list, wherever it is shown.
CARRIED_NOTE: Final = (
    "Due after the previous day's ask time, so asked in this day's batch — listed apart from the day's own."
)


# ── the day's content ──────────────────────────────────────────────────


@dataclass(slots=True)
class DigestLine:
    person: str
    email: str
    task: str
    end_user: str
    due: str
    submission_status: str
    asked: str
    mailed: str
    status: str
    reason: str
    answered: str
    managers_told: str
    task_url: str | None = None
    #: Asked in a daily batch about a task due the day before: its own list.
    carried: bool = False


@dataclass(slots=True)
class TaskLine:
    """One task due on the report's day, with everything the list holds on it."""

    person: str
    email: str
    task: str
    end_user: str
    due: str
    status: str
    submission_status: str
    current_type: str
    priority: str
    quote_no: str
    order_status: str
    remarks: str
    #: The reason given, or where its follow-up stands. Blank if never asked.
    followup: str
    task_url: str | None = None


@dataclass(slots=True)
class Digest:
    #: The last day the report covers.
    day: date
    window_start: datetime
    window_end: datetime
    zone: str
    team: str | None
    #: 1 for the end-of-day report, 7 for the weekly one.
    days: int = 1
    lines: list[DigestLine] = field(default_factory=list)
    #: Every task of the team due on the day, split by Submission Status.
    submitted: list[TaskLine] = field(default_factory=list)
    not_submitted: list[TaskLine] = field(default_factory=list)

    def count(self, status: str) -> int:
        return sum(1 for line in self.lines if line.status == STATUS_LABELS.get(status, status))

    @property
    def weekly(self) -> bool:
        return self.days > 1

    @property
    def first_day(self) -> date:
        return self.day - timedelta(days=self.days - 1)

    @property
    def kind(self) -> str:
        return "Weekly report" if self.weekly else "End of day report"

    @property
    def period(self) -> str:
        """"Monday 28 September 2026", or "22 – 28 Sep 2026" for a week."""
        if not self.weekly:
            return f"{self.day:%A %d %B %Y}"
        first = self.first_day
        head = f"{first:%d}" if (first.month, first.year) == (self.day.month, self.day.year) else f"{first:%d %b}"
        return f"{head} – {self.day:%d %b %Y}"

    @property
    def span(self) -> str:
        """How the period reads in running text: "today" or "this week"."""
        return "this week" if self.weekly else "today"

    @property
    def title(self) -> str:
        return f"{self.kind} — {self.period}"

    def people(self) -> list[tuple[str, int, int, int, int, int]]:
        """Per person: due, submitted, not submitted, reasons given, not responded."""
        table: dict[str, list[int]] = {}
        for line in self.submitted:
            table.setdefault(line.person, [0, 0, 0, 0, 0])[1] += 1
        for line in self.not_submitted:
            table.setdefault(line.person, [0, 0, 0, 0, 0])[2] += 1
        for line in self.lines:
            row = table.setdefault(line.person, [0, 0, 0, 0, 0])
            if line.status == STATUS_LABELS[FollowupStatus.ANSWERED]:
                row[3] += 1
            elif line.status == STATUS_LABELS[FollowupStatus.NO_RESPONSE]:
                row[4] += 1
        out = [(name, v[1] + v[2], v[1], v[2], v[3], v[4]) for name, v in table.items()]
        # Most not submitted first, then most not responded.
        out.sort(key=lambda r: (-r[3], -r[5], r[0].casefold()))
        return out


def _gulf(value: datetime | None) -> str:
    return value.astimezone(_GULF).strftime("%d %b %Y, %H:%M") if value else ""


async def day_tasks(
    session: AsyncSession, row: FollowupSettings, sharepoint, day: date, *, days: int = 1
) -> list[tuple[User, Any]]:
    """Every task of the team due on ``day`` (a UAE calendar day), with its holder.

    The whole team, not the trial's named people: this half of the report is
    the team's day, and the CEO reads it for the team. Read from SharePoint,
    never written.
    """
    if row.team_id is None or sharepoint is None:
        return []
    from app.followups.service import due_of
    from app.teams import service as teams_service

    first = day - timedelta(days=days - 1)
    start = datetime.combine(first, time(0, 0), tzinfo=_GULF).astimezone(UTC)
    end = start + timedelta(days=days)
    out: list[tuple[User, Any]] = []
    for user, _ in await teams_service.list_members(session, row.team_id):
        if not user.is_active:
            continue
        lookup = await sharepoint.lookup_id_for(user.email)
        if lookup is None:
            continue
        for task in await sharepoint.tasks_assigned_to(lookup, limit=500):
            due = due_of(task)
            if due is not None and start <= due < end:
                out.append((user, task))
    out.sort(key=lambda pair: due_of(pair[1]))
    return out


async def gather(
    session: AsyncSession,
    row: FollowupSettings,
    *,
    day: date,
    mark_no_response: bool,
    tasks: list[tuple[User, Any]] | None = None,
    days: int = 1,
) -> Digest:
    """The follow-ups asked in the window ending at ``day``'s closing time, and —
    at the real end of day — the unanswered ones marked not responded."""
    end = cutoff_for(row, day)
    start = cutoff_for(row, day - timedelta(days=days))
    rows = (
        await session.scalars(
            select(TaskFollowup)
            .where(TaskFollowup.created_at > start, TaskFollowup.created_at <= end)
            # A trial (no team) is the tester's alone.
            .where(TaskFollowup.team_id.is_not(None))
            .order_by(TaskFollowup.created_at)
        )
    ).all()

    if mark_no_response:
        for followup in rows:
            if followup.status == FollowupStatus.PENDING:
                followup.status = FollowupStatus.NO_RESPONSE
                followup.resolved_note = (
                    f"No reason given by the end of the day ({day:%d %b %Y}, "
                    f"{closing_time(row):%H:%M} {row.digest_timezone})."
                )
        await session.flush()

    digest = Digest(
        day=day,
        window_start=start,
        window_end=end,
        zone=row.digest_timezone or "Asia/Kolkata",
        team=row.team.name if row.team else None,
        days=days,
    )
    for f in rows:
        told = ""
        if f.status == FollowupStatus.ANSWERED:
            told = "Yes" if f.forwarded_at else "No"
        digest.lines.append(
            DigestLine(
                person=f.assignee.display_name if f.assignee else f.assignee_email,
                email=f.assignee_email,
                task=f.task_title,
                end_user=f.end_user or "",
                due=_gulf(f.due_at),
                submission_status=f.status_at_ask or "Not set",
                asked=_gulf(f.created_at),
                mailed="Yes" if f.asked_at else "No",
                status=STATUS_LABELS.get(f.status, f.status),
                reason=(f.reason or f.resolved_note or "").strip(),
                answered=_gulf(f.answered_at),
                managers_told=told,
                task_url=f.task_url,
                carried=bool(getattr(f, "carried_over", False)),
            )
        )

    # The day's tasks, each with its reason when it was asked about.
    from app.followups.service import due_of, is_finished

    by_task: dict[str, TaskFollowup] = {}
    if tasks:
        ids = {task.id for _, task in tasks}
        asked = (
            await session.scalars(
                select(TaskFollowup)
                .where(TaskFollowup.task_id.in_(ids), TaskFollowup.team_id.is_not(None))
                .order_by(TaskFollowup.created_at)
            )
        ).all()
        by_task = {f.task_id: f for f in asked}  # the latest wins
    for user, task in tasks or []:
        f = by_task.get(task.id)
        followup = ""
        if f is not None:
            followup = STATUS_LABELS.get(f.status, f.status)
            if f.reason:
                followup += f": {f.reason.strip()}"
        line = TaskLine(
            person=user.display_name,
            email=user.email,
            task=task.title or "(untitled)",
            end_user=task.end_user or "",
            due=_gulf(due_of(task)),
            status=task.status or "",
            submission_status=task.submission_status or "Not set",
            current_type=task.current_type or "",
            priority=task.priority or "",
            quote_no=task.quote_no or "",
            order_status=task.order_status or "",
            remarks=(task.remarks or "").strip(),
            followup=followup,
            task_url=task.web_url,
        )
        (digest.submitted if is_finished(task) else digest.not_submitted).append(line)
    return digest


async def recipients(session: AsyncSession, row: FollowupSettings) -> list[str]:
    """The named addresses, plus whoever holds the CEO role when that is on."""
    out: list[str] = [e.strip().lower() for e in (row.digest_recipients or []) if e.strip()]
    if row.digest_include_ceo:
        ceos = await session.scalars(
            select(User)
            .join(UserRole, UserRole.user_id == User.id)
            .join(Role, Role.id == UserRole.role_id)
            .where(Role.key == "ceo", User.is_active.is_(True))
        )
        for user in ceos.all():
            email = (user.email or "").strip().lower()
            if email and email not in out:
                out.append(email)
    return out


# ── the PDF, on the costing report's letterhead ────────────────────────


def build_pdf(digest: Digest) -> bytes:
    from reportlab.lib.enums import TA_LEFT
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.utils import ImageReader
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    from app.quoting import report_pdf as house

    page_w, page_h = landscape(A4)
    margin = house.MARGIN
    width = page_w - 2 * margin
    logo = ImageReader(str(house.LOGO)) if house.LOGO.exists() else None
    footer = ImageReader(str(house.FOOTER)) if house.FOOTER.exists() else None

    def draw(canvas, doc) -> None:
        canvas.saveState()
        y = page_h
        if logo is not None:
            canvas.drawImage(logo, margin, y - 82, width=69, height=54, mask="auto")
            canvas.setFont("Helvetica", 6.6)
            canvas.setFillColor(house.FAINT)
            canvas.drawString(margin + 3, y - 93, "w w w . h a m d a z . c o m")
        canvas.setFont("Helvetica-Bold", 9.4)
        canvas.setFillColor(house.NAVY)
        canvas.drawString(118, y - 41, house.ORG)
        canvas.setFont("Helvetica", 8.2)
        canvas.setFillColor(house.MUTED)
        for i, line in enumerate(house.ADDRESS):
            canvas.drawString(118, y - 56 - i * 12.3, line)

        canvas.setFont("Helvetica-Bold", 14)
        canvas.setFillColor(house.NAVY)
        canvas.drawRightString(page_w - margin, y - 45, digest.kind.upper())
        rule_w, rule_h = 181.0, 2.4
        canvas.setFillColor(house.CYAN)
        canvas.rect(page_w - margin - rule_w, y - 53, rule_w * 0.62, rule_h, stroke=0, fill=1)
        canvas.setFillColor(house.MAGENTA)
        canvas.rect(page_w - margin - rule_w * 0.38, y - 53, rule_w * 0.38, rule_h, stroke=0, fill=1)
        canvas.setFont("Helvetica", 8.4)
        canvas.setFillColor(house.MUTED)
        canvas.drawRightString(
            page_w - margin, y - 68, f"{digest.team or 'All teams'}  |  Overdue tasks & reasons"
        )
        canvas.setFont("Helvetica-Bold", 8.4)
        canvas.setFillColor(house.INK)
        canvas.drawRightString(page_w - margin, y - 82, digest.period)
        canvas.setStrokeColor(house.LINE)
        canvas.setLineWidth(0.8)
        canvas.line(margin, y - 100, page_w - margin, y - 100)

        if footer is not None:
            iw, ih = footer.getSize()
            fh = width * ih / iw
            canvas.drawImage(footer, margin, 17, width=width, height=fh, mask="auto")
            canvas.setFont("Helvetica", 6.5)
            canvas.setFillColor(house.FAINT)
            canvas.drawString(margin, 17 + fh + 4, f"{digest.kind}  ·  {digest.team or 'All teams'}  ·  {digest.period}")
            canvas.drawRightString(page_w - margin, 17 + fh + 4, f"Page {doc.page}")
        canvas.restoreState()

    cell = ParagraphStyle("cell", fontName="Helvetica", fontSize=7.6, leading=9.6, alignment=TA_LEFT, textColor=house.INK)
    head = ParagraphStyle("head", parent=cell, fontName="Helvetica-Bold", textColor="#ffffff")
    lead = ParagraphStyle("lead", parent=cell, fontSize=9.4, leading=13)
    section = ParagraphStyle("section", parent=lead, fontName="Helvetica-Bold", fontSize=11, textColor=house.NAVY, spaceBefore=6)

    def p(text: str, style=cell) -> Paragraph:
        return Paragraph(escape(text or "").replace("\n", "<br/>"), style)

    counts = [
        ("Due " + digest.span, len(digest.submitted) + len(digest.not_submitted)),
        ("Submitted", len(digest.submitted)),
        ("Not submitted", len(digest.not_submitted)),
        ("Reasons asked", len(digest.lines)),
        ("Reason given", digest.count(FollowupStatus.ANSWERED)),
        ("Not responded", digest.count(FollowupStatus.NO_RESPONSE)),
    ]
    tiles = Table(
        [[Paragraph(f"<font size=15><b>{n}</b></font><br/><font color='#64727f'>{escape(label)}</font>", lead) for label, n in counts]],
        colWidths=[width / len(counts)] * len(counts),
    )
    tiles.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, -1), house.PANEL),
        ("BOX", (0, 0), (-1, -1), 0.6, house.LINE),
        ("INNERGRID", (0, 0), (-1, -1), 0.6, house.LINE),
        ("TOPPADDING", (0, 0), (-1, -1), 8),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 8),
        ("LEFTPADDING", (0, 0), (-1, -1), 10),

    ]))

    headings = ["#", "Person", "Task", "Due (UAE)", "Submission", "Asked (UAE)", "Mailed", "Status", "Reason / note", "Answered (UAE)", "Managers told"]
    widths = [18, 64, 150, 62, 52, 62, 34, 58, 170, 62, 46]
    scale = width / sum(widths)

    def reasons_table(lines: list[DigestLine]) -> Table:
        data: list[list[Any]] = [[p(h, head) for h in headings]]
        for i, line in enumerate(lines, start=1):
            task = line.task + (f"\n{line.end_user}" if line.end_user else "")
            data.append([
                p(str(i)), p(line.person), p(task), p(line.due), p(line.submission_status),
                p(line.asked), p(line.mailed), p(line.status), p(line.reason),
                p(line.answered), p(line.managers_told),
            ])
        table = Table(data, colWidths=[w * scale for w in widths], repeatRows=1)
        style = [
            ("BACKGROUND", (0, 0), (-1, 0), house.NAVY),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("LINEBELOW", (0, 0), (-1, -1), 0.5, house.LINE),
            ("TOPPADDING", (0, 0), (-1, -1), 4),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
            ("LEFTPADDING", (0, 0), (-1, -1), 4),
            ("RIGHTPADDING", (0, 0), (-1, -1), 4),
        ]
        for r, line in enumerate(lines, start=1):
            if line.status == STATUS_LABELS[FollowupStatus.NO_RESPONSE]:
                style.append(("BACKGROUND", (0, r), (-1, r), house.LOSS_RED))
            elif r % 2 == 0:
                style.append(("BACKGROUND", (0, r), (-1, r), house.ROW))
        table.setStyle(TableStyle(style))
        return table

    own = [line for line in digest.lines if not line.carried]
    carried = [line for line in digest.lines if line.carried]

    def task_table(lines: list[TaskLine], shade) -> Table:
        heads = ["#", "Person", "Task", "Due (UAE)", "Status", "Submission", "Type", "Priority", "Zoho quote", "Order", "Remarks", "Reason / follow-up"]
        cols = [16, 56, 130, 56, 46, 50, 34, 38, 50, 40, 130, 110]
        k = width / sum(cols)
        rows: list[list[Any]] = [[p(h, head) for h in heads]]
        for n, line in enumerate(lines, start=1):
            task = line.task + (f"\n{line.end_user}" if line.end_user else "")
            remarks = line.remarks if len(line.remarks) <= 300 else line.remarks[:300] + "…"
            rows.append([
                p(str(n)), p(line.person), p(task), p(line.due), p(line.status),
                p(line.submission_status), p(line.current_type), p(line.priority),
                p(line.quote_no), p(line.order_status), p(remarks), p(line.followup),
            ])
        t = Table(rows, colWidths=[c * k for c in cols], repeatRows=1)
        st = [
            ("BACKGROUND", (0, 0), (-1, 0), house.NAVY),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("LINEBELOW", (0, 0), (-1, -1), 0.5, house.LINE),
            ("TOPPADDING", (0, 0), (-1, -1), 4),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
            ("LEFTPADDING", (0, 0), (-1, -1), 4),
            ("RIGHTPADDING", (0, 0), (-1, -1), 4),
        ]
        for r in range(1, len(rows)):
            if shade is not None:
                st.append(("BACKGROUND", (0, r), (-1, r), shade))
            elif r % 2 == 0:
                st.append(("BACKGROUND", (0, r), (-1, r), house.ROW))
        t.setStyle(TableStyle(st))
        return t

    when = (
        f"from {digest.first_day:%d %B} to {digest.day:%d %B %Y}" if digest.weekly
        else f"on {digest.day:%d %B %Y}"
    )
    story: list[Any] = [
        Paragraph(
            f"{digest.team or 'The team'} tasks due {when} (UAE time), "
            f"submitted and not, and every reason asked for between "
            f"{_local(digest.window_start, digest.zone)} and {_local(digest.window_end, digest.zone)}. "
            f"Anybody who had not answered by the day's closing time is marked <b>Not responded</b>.",
            lead,
        ),
        Spacer(1, 8),
        tiles,
        Spacer(1, 10),
    ]
    if digest.weekly:
        # The week at a glance, one row a person, before the task lists.
        people = digest.people()
        rows = [[p(h, head) for h in ("Person", "Due", "Submitted", "Not submitted", "Reasons given", "Not responded")]]
        rows += [[p(name), p(str(a)), p(str(b)), p(str(c)), p(str(d)), p(str(e))] for name, a, b, c, d, e in people]
        people_table = Table(rows, colWidths=[width * 0.3] + [width * 0.14] * 5, repeatRows=1)
        people_table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), house.NAVY),
            ("LINEBELOW", (0, 0), (-1, -1), 0.5, house.LINE),
            ("TOPPADDING", (0, 0), (-1, -1), 4),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
            ("LEFTPADDING", (0, 0), (-1, -1), 6),
        ]))
        story += [
            Paragraph("By person", section),
            Spacer(1, 4),
            people_table if people else Paragraph("Nothing was due this week.", lead),
            Spacer(1, 10),
        ]
    story += [
        Paragraph(f"Not submitted ({len(digest.not_submitted)})", section),
        Spacer(1, 4),
        task_table(digest.not_submitted, house.LOSS_RED) if digest.not_submitted
        else Paragraph(f"Every task due {digest.span} was submitted.", lead),
        Spacer(1, 10),
        Paragraph(f"Submitted ({len(digest.submitted)})", section),
        Spacer(1, 4),
        task_table(digest.submitted, None) if digest.submitted
        else Paragraph(f"Nothing due {digest.span} was submitted.", lead),
        Spacer(1, 10),
        Paragraph(f"Reasons asked for ({len(own)})", section),
        Spacer(1, 4),
    ]
    story.append(
        reasons_table(own) if own
        else Paragraph(f"Nobody was asked for a reason {digest.span}.", lead)
    )
    if carried:
        story += [
            Spacer(1, 10),
            Paragraph(f"Carried over from the previous day ({len(carried)})", section),
            Paragraph(CARRIED_NOTE, lead),
            Spacer(1, 4),
            reasons_table(carried),
        ]

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer, pagesize=(page_w, page_h), leftMargin=margin, rightMargin=margin,
        topMargin=house.TOP, bottomMargin=house.BOTTOM, title=digest.title, author="Hamdaz ERP",
    )
    doc.build(story, onFirstPage=draw, onLaterPages=draw)
    return buffer.getvalue()


#: How a timezone is named in the report's prose.
_ZONE_NAMES = {"Asia/Kolkata": "India time", "Asia/Calcutta": "India time", "Asia/Dubai": "UAE time"}


def _local(value: datetime, zone: str) -> str:
    name = _ZONE_NAMES.get(zone, zone.split("/")[-1] + " time")
    return value.astimezone(_zone(zone)).strftime(f"%d %b %H:%M ({name})")


# ── the workbook ───────────────────────────────────────────────────────


def build_xlsx(digest: Digest) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    navy = "0E5E80"
    wb = Workbook()
    ws = wb.active
    ws.title = "Reasons"
    thin = Side(style="thin", color="E2E9EF")

    ws["A1"] = "HAMDAZTECH TECHNOLOGY SERVICES - L.L.C"
    ws["A1"].font = Font(bold=True, size=11, color=navy)
    ws["A2"] = f"Reasons asked for — {digest.kind.lower()}"
    ws["A2"].font = Font(bold=True, size=14, color=navy)
    ws["A3"] = f"{digest.period}  ·  {digest.team or 'All teams'}"
    ws["A3"].font = Font(size=10, color="64727F")
    summary = (
        f"Due {digest.span} {len(digest.submitted) + len(digest.not_submitted)}  ·  "
        f"Submitted {len(digest.submitted)}  ·  Not submitted {len(digest.not_submitted)}  ·  "
        f"Reasons asked {len(digest.lines)}  ·  Reason given {digest.count(FollowupStatus.ANSWERED)}  ·  "
        f"Not responded {digest.count(FollowupStatus.NO_RESPONSE)}  ·  "
        f"False positive {digest.count(FollowupStatus.FALSE_POSITIVE)}  ·  "
        f"Closed {digest.count(FollowupStatus.RESOLVED)}"
    )
    ws["A4"] = summary
    ws["A4"].font = Font(bold=True, size=10)

    headings = [
        "#", "Person", "Email", "Task", "End user", "Due (UAE)", "Submission status",
        "Asked (UAE)", "Mailed", "Status", "Reason / note", "Answered (UAE)", "Managers told",
        "SharePoint link", "List",
    ]
    start = 6
    for col, heading in enumerate(headings, start=1):
        cell = ws.cell(row=start, column=col, value=heading)
        cell.font = Font(bold=True, color="FFFFFF", size=10)
        cell.fill = PatternFill("solid", fgColor=navy)
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    red = PatternFill("solid", fgColor="FDE2E1")
    ordered = [x for x in digest.lines if not x.carried] + [x for x in digest.lines if x.carried]
    for i, line in enumerate(ordered, start=1):
        values = [
            i, line.person, line.email, line.task, line.end_user, line.due,
            line.submission_status, line.asked, line.mailed, line.status, line.reason,
            line.answered, line.managers_told, line.task_url or "",
            "Carried over from the previous day" if line.carried else "The day's own",
        ]
        for col, value in enumerate(values, start=1):
            cell = ws.cell(row=start + i, column=col, value=value)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = Border(bottom=thin)
            if line.status == STATUS_LABELS[FollowupStatus.NO_RESPONSE]:
                cell.fill = red
    for col, w in enumerate([5, 18, 24, 44, 24, 17, 16, 17, 8, 15, 60, 17, 10, 40, 30], start=1):
        ws.column_dimensions[ws.cell(row=start, column=col).column_letter].width = w
    ws.freeze_panes = ws.cell(row=start + 1, column=1)
    ws.auto_filter.ref = f"A{start}:{ws.cell(row=start, column=len(headings)).column_letter}{start + max(1, len(digest.lines))}"

    task_heads = [
        "#", "Person", "Email", "Task", "End user", "Due (UAE)", "Status", "Submission status",
        "Current type", "Priority", "Zoho quote no", "Order status", "Remarks",
        "Reason / follow-up", "SharePoint link",
    ]
    task_widths = [5, 18, 24, 44, 24, 17, 14, 16, 12, 10, 16, 14, 60, 50, 40]
    for title, lines, shade in (
        ("Not submitted", digest.not_submitted, red),
        ("Submitted", digest.submitted, None),
    ):
        sheet = wb.create_sheet(title)
        sheet["A1"] = "HAMDAZTECH TECHNOLOGY SERVICES - L.L.C"
        sheet["A1"].font = Font(bold=True, size=11, color=navy)
        sheet["A2"] = f"{title} — due {digest.period} (UAE)"
        sheet["A2"].font = Font(bold=True, size=14, color=navy)
        sheet["A3"] = f"{digest.team or 'All teams'}  ·  {len(lines)} task(s)"
        sheet["A3"].font = Font(size=10, color="64727F")
        top = 5
        for col, heading in enumerate(task_heads, start=1):
            cell = sheet.cell(row=top, column=col, value=heading)
            cell.font = Font(bold=True, color="FFFFFF", size=10)
            cell.fill = PatternFill("solid", fgColor=navy)
            cell.alignment = Alignment(vertical="center", wrap_text=True)
        for n, line in enumerate(lines, start=1):
            values = [
                n, line.person, line.email, line.task, line.end_user, line.due, line.status,
                line.submission_status, line.current_type, line.priority, line.quote_no,
                line.order_status, line.remarks, line.followup, line.task_url or "",
            ]
            for col, value in enumerate(values, start=1):
                cell = sheet.cell(row=top + n, column=col, value=value)
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                cell.border = Border(bottom=thin)
                if shade is not None:
                    cell.fill = shade
        for col, w in enumerate(task_widths, start=1):
            sheet.column_dimensions[sheet.cell(row=top, column=col).column_letter].width = w
        sheet.freeze_panes = sheet.cell(row=top + 1, column=1)
    # The day's picture first; the reasons behind it after.
    wb.move_sheet("Reasons", offset=2)

    if digest.weekly:
        people = wb.create_sheet("By person", 0)
        people["A1"] = "HAMDAZTECH TECHNOLOGY SERVICES - L.L.C"
        people["A1"].font = Font(bold=True, size=11, color=navy)
        people["A2"] = f"Weekly report — {digest.period}"
        people["A2"].font = Font(bold=True, size=14, color=navy)
        people["A3"] = digest.team or "All teams"
        people["A3"].font = Font(size=10, color="64727F")
        heads = ["Person", "Due", "Submitted", "Not submitted", "Reasons given", "Not responded"]
        for col, heading in enumerate(heads, start=1):
            cell = people.cell(row=5, column=col, value=heading)
            cell.font = Font(bold=True, color="FFFFFF", size=10)
            cell.fill = PatternFill("solid", fgColor=navy)
        for n, values in enumerate(digest.people(), start=1):
            for col, value in enumerate(values, start=1):
                cell = people.cell(row=5 + n, column=col, value=value)
                cell.border = Border(bottom=thin)
        for col, w in enumerate([28, 10, 12, 14, 14, 14], start=1):
            people.column_dimensions[people.cell(row=5, column=col).column_letter].width = w

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


# ── the mail ───────────────────────────────────────────────────────────


def _follow_up_state(text: str) -> str:
    """The follow-up column, without the reason itself — that is in the attachment."""
    if not text:
        return "Not asked yet"
    return text.split(":", 1)[0]


def mail_html(digest: Digest, link: str) -> str:
    """The body: the day's numbers, then the tasks not submitted, in one table.

    Only what somebody decides on from an inbox — who, what, when, and where
    the follow-up stands. The reasons in full, the submitted list and every
    other column are in the PDF and the workbook attached.
    """
    from app.followups import mailer as m

    body = m.heading(digest.title, eyebrow=f"Proposals · {digest.team or 'All teams'}")
    body += m.counts([
        ("Due " + digest.span, len(digest.submitted) + len(digest.not_submitted), False),
        ("Submitted", len(digest.submitted), False),
        ("Not submitted", len(digest.not_submitted), True),
        ("Not responded", digest.count(FollowupStatus.NO_RESPONSE), True),
    ])
    if digest.weekly and digest.people():
        body += (
            f"<p style='margin:20px 0 8px;font-weight:600'>By person</p>"
            + m.grid(
                ["Person", "Due", "Submitted", "Not submitted", "Not responded"],
                [[name, str(a), str(b), str(c), str(e)] for name, a, b, c, _, e in digest.people()],
            )
        )
    if digest.not_submitted:
        # A week can hold dozens; the mail shows the first and says where the rest are.
        shown = digest.not_submitted[:25]
        rows = [
            [str(i), t.person, t.task, t.due if digest.weekly else t.due.split(", ")[-1], _follow_up_state(t.followup)]
            for i, t in enumerate(shown, start=1)
        ]
        colours = [None, None, None, None, None]
        more = len(digest.not_submitted) - len(shown)
        body += (
            f"<p style='margin:20px 0 8px;font-weight:600'>Not submitted</p>"
            + m.grid(["#", "Person", "Task", "Due (UAE)", "Follow-up"], rows, colours=colours)
            + (f"<p style='margin:6px 0 0;font-size:12.5px;color:#5f6b77'>And {more} more in the "
               f"attached report.</p>" if more > 0 else "")
        )
    else:
        body += f"<p style='margin:20px 0 0'>Every task due {digest.span} was submitted.</p>"
    carried_n = sum(1 for line in digest.lines if line.carried)
    if carried_n:
        body += (
            f"<p style='margin:20px 0 0'><b>{carried_n} carried over from the previous day</b> "
            f"\u2014 {escape(CARRIED_NOTE[0].lower() + CARRIED_NOTE[1:])} They are in their own "
            f"section of the attached report.</p>"
        )
    missing = [line for line in digest.lines if line.status == STATUS_LABELS[FollowupStatus.NO_RESPONSE]]
    if missing:
        body += (
            f"<p style='margin:20px 0 8px;font-weight:600'>No reason given by the closing time</p>"
            + m.grid(
                ["Person", "Task"],
                [[line.person, line.task] for line in missing[:25]],
            )
        )
    body += (
        f"<p style='margin:18px 0 0;font-size:12.5px;color:#5f6b77'>The full report — every task "
        f"with its details, and every reason given — is attached as PDF and Excel.</p>"
        f"<p style='margin:16px 0 0'>{m.button(link, 'View Report')}</p>"
    )
    return m.page(body)


def attachments(digest: Digest, formats: list[str]) -> list[Attachment]:
    stem = (
        f"weekly-report-{digest.first_day:%Y-%m-%d}-to-{digest.day:%Y-%m-%d}"
        if digest.weekly
        else f"end-of-day-report-{digest.day:%Y-%m-%d}"
    )
    out: list[Attachment] = []
    if "pdf" in formats:
        out.append(Attachment(f"{stem}.pdf", build_pdf(digest), PDF_TYPE))
    if "xlsx" in formats:
        out.append(Attachment(f"{stem}.xlsx", build_xlsx(digest), XLSX_TYPE))
    return out


async def send(
    session: AsyncSession,
    row: FollowupSettings,
    *,
    mailer,
    link: str,
    now: datetime,
    sharepoint=None,
    day: date | None = None,
    weekly: bool = False,
    preview: bool = False,
) -> dict[str, Any]:
    """Gather and mail one report — the day's, or the week to ``day``.

    The real end-of-day send marks the unanswered as not responded and
    records the day; the real weekly send records the day; a preview does
    neither, so trying it out never changes what the real one will say.
    """
    day = day or local_today(row, now)
    days = 7 if weekly else 1
    mark_no_response = not weekly and not preview
    try:
        tasks = await day_tasks(session, row, sharepoint, day, days=days)
    except Exception as exc:  # noqa: BLE001 - the reasons still go, said so
        logger.warning("report: the tasks could not be read: %s", exc)
        tasks = []
    digest = await gather(
        session, row, day=day, mark_no_response=mark_no_response, tasks=tasks, days=days
    )
    to = await recipients(session, row)
    result: dict[str, Any] = {
        "day": day.isoformat(),
        "period": "week" if weekly else "day",
        "lines": len(digest.lines),
        "submitted": len(digest.submitted),
        "not_submitted": len(digest.not_submitted),
        "not_responded": digest.count(FollowupStatus.NO_RESPONSE),
        "recipients": to,
        "sent": False,
        "error": None,
    }
    if not to:
        result["error"] = "Nobody to send the report to — add an address in the settings."
    else:
        from app.followups import mailer as m

        sender = (row.digest_sender_email or to[0]).strip()
        try:
            await mailer.send(
                sender=sender,
                recipients=to,
                subject=f"{digest.kind.title()} — {digest.period} — "
                f"{len(digest.submitted)} Submitted, {len(digest.not_submitted)} Not Submitted",
                html=mail_html(digest, link),
                attachments=attachments(digest, list(row.digest_formats or ["pdf", "xlsx"]))
                + m.logo_attachment(),
            )
            result["sent"] = True
        except Exception as exc:  # noqa: BLE001 - recorded, and tried again tomorrow
            result["error"] = f"{type(exc).__name__}: {exc}"[:2000]
            logger.warning("%s for %s not sent: %s", digest.kind, day, exc)
    if not preview:
        # Recorded even when the mail failed, so the send is not repeated on
        # every tick — the error on the settings says it needs resending.
        if weekly:
            row.weekly_last_sent_on = day
        else:
            row.digest_last_sent_on = day
    row.digest_last_error = result["error"]
    await session.flush()
    return result
